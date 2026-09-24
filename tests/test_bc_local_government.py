"""Evidence checks for the real-data B.C. local-government workbook."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from workbook_io import ROOT


PROJECT = ROOT / "examples" / "bc-local-government"
MUNICIPAL = PROJECT / "data" / "municipal_finance.csv"
ASSETS = PROJECT / "data" / "asset_categories_2024.csv"
WORKBOOK = PROJECT / "BC_Local_Government_Finance_Model.xlsx"
EXPECTED_SHEETS = (
    "Dashboard",
    "Forecast",
    "Capital plan",
    "Assumptions",
    "Checks",
    "Municipal data",
    "Asset data",
    "ReadMe",
)


@pytest.fixture(scope="module")
def municipal() -> pd.DataFrame:
    return pd.read_csv(MUNICIPAL)


@pytest.fixture(scope="module")
def assets() -> pd.DataFrame:
    return pd.read_csv(ASSETS)


@pytest.fixture(scope="module")
def workbook() -> openpyxl.Workbook:
    return openpyxl.load_workbook(WORKBOOK, data_only=False, read_only=True)


def test_committed_summary_matches_the_release():
    summary = json.loads((PROJECT / "data" / "summary.json").read_text(encoding="utf-8"))
    assert summary["years"] == [2019, 2020, 2021, 2022, 2023, 2024]
    assert summary["municipality_year_rows"] == 968
    assert summary["municipalities_2024"] == 161
    assert summary["asset_rows_2024"] == 1932
    assert summary["source_files"] == 66


def test_municipality_year_key_is_unique(municipal: pd.DataFrame):
    assert len(municipal) == 968
    assert not municipal["key"].duplicated().any()
    assert "Grand Totals" not in set(municipal["municipality_name"])


@pytest.mark.parametrize("year", range(2019, 2025))
def test_each_year_has_the_expected_reporting_entities(municipal: pd.DataFrame, year: int):
    expected = 162 if year < 2021 else 161
    assert len(municipal.loc[municipal["year"].eq(year)]) == expected


def test_entity_history_is_continuous(municipal: pd.DataFrame):
    observed = municipal.groupby("municipality_id")["year"].agg(list).to_dict()
    jumbo = observed.pop("jumbo-glacier-v")
    assert jumbo == [2019, 2020]
    assert set(map(tuple, observed.values())) == {(2019, 2020, 2021, 2022, 2023, 2024)}


@pytest.mark.parametrize(
    ("total", "components"),
    [
        (
            "total_revenue",
            (
                "taxation_revenue", "sale_of_services", "federal_transfers",
                "provincial_transfers", "regional_transfers", "investment_income",
                "government_business_income", "developer_contributions",
                "gain_on_sale_assets", "other_revenue",
            ),
        ),
        (
            "total_expenses",
            (
                "general_government_expense", "protective_services_expense",
                "solid_waste_expense", "health_social_housing_expense",
                "development_services_expense", "transportation_expense",
                "parks_recreation_expense", "water_expense", "sewer_expense",
                "other_services_expense", "amortization",
                "asset_retirement_accretion", "loss_on_disposition",
                "other_adjustments_expense",
            ),
        ),
    ],
)
def test_published_components_reconcile(municipal: pd.DataFrame, total: str, components: tuple[str, ...]):
    difference = municipal[list(components)].fillna(0).sum(axis=1) - municipal[total]
    assert difference.abs().max() == 0


def test_asset_key_is_unique(assets: pd.DataFrame):
    assert len(assets) == 1932
    assert not assets["key"].duplicated().any()
    assert assets["municipality_id"].nunique() == 161


def test_asset_categories_reconcile_to_total_tca(assets: pd.DataFrame):
    detail = assets.loc[~assets["asset_category"].eq("Total TCA")]
    total = assets.loc[assets["asset_category"].eq("Total TCA")].set_index("municipality_id")
    grouped = detail.groupby("municipality_id")[["historical_cost", "accumulated_amortization", "net_book_value", "asset_additions"]].sum(min_count=1)
    pd.testing.assert_frame_equal(grouped, total[grouped.columns], check_dtype=False)


@pytest.mark.parametrize("sheet", EXPECTED_SHEETS)
def test_workbook_contains_each_required_sheet(workbook: openpyxl.Workbook, sheet: str):
    assert sheet in workbook.sheetnames


@pytest.mark.parametrize("sheet", ("Dashboard", "Forecast", "Capital plan", "Checks"))
def test_decision_sheets_contain_live_formulas(workbook: openpyxl.Workbook, sheet: str):
    formulas = [cell.value for row in workbook[sheet].iter_rows() for cell in row if cell.data_type == "f"]
    assert len(formulas) >= 5


def test_workbook_contains_no_stored_formula_errors(workbook: openpyxl.Workbook):
    errors = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    assert not errors


def test_three_planning_cases_are_declared(workbook: openpyxl.Workbook):
    values = {cell.value for row in workbook["Assumptions"].iter_rows() for cell in row}
    assert {"Base", "Upside", "Downside"} <= values


def test_source_manifest_fingerprints_all_official_workbooks():
    manifest = json.loads((PROJECT / "data" / "source_manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["files"]) == 66
    assert len({item["filename"] for item in manifest["files"]}) == 66
    assert all(len(item["sha256"]) == 64 for item in manifest["files"])
    assert all(item["url"].startswith("https://www2.gov.bc.ca/") for item in manifest["files"])
