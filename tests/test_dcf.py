"""The lululemon DCF workbook, held to its Python twin and to the filings.

valuation/build_workbook.py writes every calculation as a formula and has
LibreOffice recalculate the file; valuation/dcf_model.py computes the same
valuation in Python from the same three input files. These tests compare them
figure by figure, check the filed figures against SEC EDGAR's own arithmetic,
and hold the README to the numbers.
"""
from __future__ import annotations

import json
import sys
import warnings

import openpyxl
import pytest

from workbook_io import ROOT

VAL = ROOT / "valuation"
sys.path.insert(0, str(VAL))
import dcf_model as twin  # noqa: E402

WORKBOOK = VAL / "Lululemon_DCF_FY2025.xlsx"
CELLS = json.loads((VAL / "cell_map.json").read_text(encoding="utf-8"))
SCENARIOS = ["bear", "base", "bull"]


def _load(data_only: bool):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(WORKBOOK, data_only=data_only)


@pytest.fixture(scope="module")
def wb():
    return _load(True)


@pytest.fixture(scope="module")
def wbf():
    return _load(False)


@pytest.fixture(scope="module")
def ref():
    return twin.run()


def cell(wb, key):
    sheet, ref_ = CELLS[key].split("!")
    return wb[sheet][ref_].value


VALUE_KEYS = ["enterprise_value", "equity_value", "value_per_share", "terminal_value", "pv_terminal", "pv_fcff",
              "terminal_share_of_ev", "implied_terminal_ev_ebitda", "ev_to_fy2025_ebit", "cash"]


@pytest.mark.parametrize("scenario", SCENARIOS)
@pytest.mark.parametrize("key", VALUE_KEYS)
def test_every_valuation_figure_matches_the_twin(wb, ref, scenario, key):
    assert cell(wb, f"{scenario}_{key}") == pytest.approx(ref["scenarios"][scenario]["value"][key], rel=1e-9, abs=1e-6)


@pytest.mark.parametrize("scenario", SCENARIOS)
def test_every_forecast_year_matches_the_twin(wb, ref, scenario):
    for row in ref["scenarios"][scenario]["forecast"]:
        y = row["year"]
        for key in ("revenue", "ebit", "nopat", "fcff"):
            assert cell(wb, f"{scenario}_{key}_{y}") == pytest.approx(row[key], rel=1e-9), (y, key)


def test_the_timing_and_cost_of_capital(wb, ref):
    for y, t in zip(twin.A["explicit_years"], ref["timing"], strict=True):
        assert cell(wb, f"t_{y}") == pytest.approx(t, abs=1e-9)
    assert cell(wb, "base_wacc") == pytest.approx(ref["cost_of_capital"]["wacc"], abs=1e-12)
    assert ref["cost_of_capital"]["wacc"] == pytest.approx(0.0420 + 0.79 * 0.0446)


def test_the_sensitivity_grid_is_the_twin_at_every_point(wb, ref):
    for i, row in enumerate(ref["sensitivity"]):
        for j, v in enumerate(row):
            assert cell(wb, f"sens_{i}_{j}") == pytest.approx(v, rel=1e-9), (i, j)
    assert cell(wb, "beta_alt_value_per_share") == pytest.approx(ref["beta_alternative"]["value_per_share"], rel=1e-9)


def test_history_matches_the_twin(wb, ref):
    names = {"revenue": "revenue", "ebit": "ebit", "ebit_margin": "ebit_margin", "tax_rate": "tax_rate",
             "nwc_pct": "nwc_pct", "da_pct": "da_pct", "capex_pct": "capex_pct", "other_ops": "other_operating_items"}
    for y, values in ref["history"].items():
        for key, twin_key in names.items():
            assert cell(wb, f"hist_{key}_{y}") == pytest.approx(values[twin_key], rel=1e-9), (y, key)


def test_every_check_in_the_workbook_passes(wb):
    sheet, rng = CELLS["checks_range"].split("!")
    statuses = [c[0].value for c in wb[sheet][rng]]
    assert statuses and set(statuses) == {"PASS"}
    assert cell(wb, "checks_passing") == int(CELLS["checks_total"]) == len(statuses)


def test_no_cell_holds_an_error(wb):
    errors = [f"{ws.title}!{c.coordinate}" for ws in wb.worksheets for r in ws.iter_rows() for c in r
              if isinstance(c.value, str) and c.value.startswith("#")]
    assert errors == []


@pytest.mark.parametrize("sheet", ["Summary", "DCF_Base", "DCF_Bear", "DCF_Bull", "Sensitivity", "History", "Checks"])
def test_no_calculation_sheet_holds_a_typed_number(wbf, sheet):
    typed = [c.coordinate for r in wbf[sheet].iter_rows() for c in r
             if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
    assert typed == [], f"{sheet} has typed numbers at {typed[:5]}"


def test_the_data_sheet_is_the_filed_figures(wbf):
    ws = wbf["Data_EDGAR"]
    found = {ws.cell(row=r, column=1).value: r for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=1).value}
    for key, f in twin.FACTS["facts"].items():
        r = found[key]
        assert ws.cell(row=r, column=2).value == f["tag"]
        for j, y in enumerate(twin.YEARS):
            assert ws.cell(row=r, column=3 + j).value == f["values"][y]["value"], (key, y)
            assert ws.cell(row=r, column=8 + j).value == f["values"][y]["accession"], (key, y)


def test_the_filed_figures_obey_their_own_arithmetic():
    for y in twin.YEARS:
        f = twin.fact
        assert f("gross_profit", y) == f("revenue", y) - f("cost_of_goods_sold", y), y
        assert f("net_income", y) == f("pretax_income", y) - f("income_tax", y), y
    for key, item in twin.FACTS["facts"].items():
        for y, v in item["values"].items():
            assert v["accession"].startswith("0001397187-"), (key, y)


def test_the_valuation_date_is_the_10k_filing_date():
    filed = {v["filed"] for f in twin.FACTS["facts"].values() for v in f["values"].values()
             if v["accession"] == "0001397187-26-000020"}
    assert filed == {twin.MARKET["valuation_date"]}
    for key in ("risk_free_rate", "equity_risk_premium", "beta"):
        assert twin.MARKET[key]["source"], key


def test_the_scenarios_tell_the_story_the_readme_tells(ref):
    v = {s: ref["scenarios"][s]["value"]["value_per_share"] for s in SCENARIOS}
    assert v["bear"] < v["base"] < v["bull"]
    assert 0.7 < ref["scenarios"]["base"]["value"]["terminal_share_of_ev"] < 0.85


def test_the_readme_states_the_numbers_the_model_computes(ref):
    text = (VAL / "README.md").read_text(encoding="utf-8")
    for s in SCENARIOS:
        assert f"${ref['scenarios'][s]['value']['value_per_share']:,.0f}" in text, s
    base = ref["scenarios"]["base"]["value"]
    assert f"{ref['cost_of_capital']['cost_of_equity']:.2%}" in text
    assert f"{base['terminal_share_of_ev']:.0%}" in text
    assert f"{base['implied_terminal_ev_ebitda']:.1f}x" in text
    grid = [v for row in ref["sensitivity"] for v in row]
    assert f"${min(grid):,.0f}" in text and f"${max(grid):,.0f}" in text
    assert f"${ref['beta_alternative']['value_per_share']:,.0f}" in text
