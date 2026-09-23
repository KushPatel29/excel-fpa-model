"""The Excel techniques the README claims, proved from the saved file.

"Built with Power Query and a Power Pivot data model" is easy to claim and easy
to fake with pasted values. Each test opens the part of the .xlsx package that
only exists if the feature is really there.
"""
from __future__ import annotations

import re
from datetime import datetime

import pytest
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

import layout as L
from workbook_io import ROOT, parts, power_query_m, sheet_xml

ERRORS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!",
          "#GETTING_DATA", "#FIELD!", "#BUSY!")


@pytest.fixture(scope="module")
def pkg():
    return parts()


def test_sheets_in_reading_order(values):
    assert values.sheetnames == L.SHEETS


@pytest.mark.parametrize("sheet", L.SHEETS)
def test_no_cell_holds_an_error(values, sheet):
    bad = [(c.coordinate, c.value) for row in values[sheet].iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith(ERRORS)]
    assert not bad, f"{sheet}: {bad[:5]}"


# ------------------------------------------------------------- Power Query
def test_four_power_query_queries(pkg):
    m = power_query_m(pkg)
    for name in ("Monthly", "Weather", "Calendar", "Class"):
        assert re.search(rf"shared #?\"?{name}\"? =", m), name


def test_eia_query_unpivots_and_pivots(pkg):
    m = power_query_m(pkg)
    monthly = m[m.index("shared Monthly"):m.index("shared Weather")]
    for step in ("Table.UnpivotOtherColumns", "Text.BeforeDelimiter", "Table.Pivot"):
        assert step in monthly, step


def test_noaa_query_parses_fixed_width_text(pkg):
    m = power_query_m(pkg)
    weather = m[m.index("shared Weather"):]
    assert "Splitter.SplitTextByPositions" in weather and "-9999" in weather


def test_query_connections(pkg):
    xml = pkg["xl/connections.xml"].decode()
    for name in ("Monthly", "Weather", "Calendar", "Class"):
        assert f'name="Query - {name}"' in xml, name


# ------------------------------------------------------- data model and DAX
def test_power_pivot_data_model_is_embedded(pkg):
    assert len(pkg["xl/model/item.data"]) > 50_000


def test_pivot_table_reads_dax_measures(pkg):
    cache = next(v.decode() for k, v in pkg.items() if k.startswith("xl/pivotCache/pivotCacheDefinition"))
    assert 'type="external"' in cache
    for measure in ("Retail Revenue", "Revenue YoY %", "Retail MWh", "MWh YoY %", "Price c/kWh", "Avg Customers"):
        assert f'uniqueName="[Measures].[{measure}]"' in cache, measure


def test_slicers_on_the_data_model(pkg):
    caches = " ".join(v.decode() for k, v in pkg.items() if k.startswith("xl/slicerCaches/"))
    for field in ("[Calendar].[year]", "[Calendar].[quarter]"):
        assert field in caches, field


def test_cube_formulas(formulas):
    cube = [c.value for row in formulas["Explore"].iter_rows() for c in row
            if isinstance(c.value, str) and "CUBEVALUE" in c.value]
    assert len(cube) >= 20


# --------------------------------------------------------- what-if analysis
def test_data_tables(pkg, formulas):
    scenarios = re.findall(r'<f t="dataTable"[^>]*>', sheet_xml(pkg, formulas, "Scenarios"))
    plan = re.findall(r'<f t="dataTable"[^>]*>', sheet_xml(pkg, formulas, "Plan"))
    assert len(scenarios) == 7, "weather, the rate x MW grid, and five tornado levers"
    assert any('dt2D="1"' in t for t in scenarios)
    assert len(plan) == 1 and 'r1="C4"' in plan[0], "the backtest re-runs the plan year"


def test_validated_levers(formulas):
    dv = [d for d in formulas["Scenarios"].data_validations.dataValidation if "C4" in str(d.sqref)]
    assert dv and dv[0].type == "list" and dv[0].formula1 == "ScenarioList"


# ------------------------------------------------------- modern formulas
def test_named_lambda_functions(formulas):
    for name in ("SAFEDIV", "ISCLOSED", "CPK", "TRENDYEARS", "MONEY", "SIGNMONEY"):
        assert formulas.defined_names[name].attr_text.startswith("_xlfn.LAMBDA("), name


def _text(formulas, sheet):
    return " ".join(str(c.value.text if hasattr(c.value, "text") else c.value)
                    for row in formulas[sheet].iter_rows() for c in row if c.value is not None)


@pytest.mark.parametrize("function,sheet", [
    ("LINEST", "Weather"), ("FILTER", "Weather"), ("HSTACK", "Weather"), ("XLOOKUP", "Forecast"),
    ("FORECAST.ETS", "Forecast"), ("SORTBY", "Scenarios"), ("LET", "Dashboard"), ("SUMIFS", "PnL"),
    ("RANK.EQ", "Peers"), ("EDATE", "Plan"),
])
def test_modern_functions_in_use(formulas, function, sheet):
    assert function in _text(formulas, sheet), f"{function} not found on {sheet}"


# ------------------------------------------------------------- charts
def test_waterfall_charts(pkg):
    chartex = [v.decode() for k, v in pkg.items() if k.startswith("xl/charts/chartEx")]
    assert sum('layoutId="waterfall"' in c for c in chartex) >= 4


def test_regular_charts(pkg):
    assert len([k for k in pkg if re.fullmatch(r"xl/charts/chart\d+\.xml", k)]) >= 5


def test_sparklines(pkg, formulas):
    assert "sparklineGroup" in sheet_xml(pkg, formulas, "Dashboard")


# ------------------------------------------------- no numbers typed by hand
def excel_owned(ws) -> set[str]:
    """Cells whose numbers Excel wrote itself: dynamic-array spills and PivotTable output."""
    def cells(ref):
        c1, r1, c2, r2 = range_boundaries(ref)
        return {f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)}

    owned = set()
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, ArrayFormula) and c.value.ref:
                owned |= cells(c.value.ref)
    for pt in getattr(ws, "_pivots", []):
        owned |= cells(pt.location.ref)
    return owned


def typed_numbers(ws, allowed=frozenset()):
    owned = excel_owned(ws) | set(allowed)
    return [c.coordinate for row in ws.iter_rows() for c in row
            if isinstance(c.value, (int, float, datetime)) and not isinstance(c.value, bool)
            and c.coordinate not in owned]


@pytest.mark.parametrize("sheet", ["Dashboard", "PnL", "PVM", "Monthly", "Weather", "Forecast", "Peers"])
def test_calculation_sheets_hold_no_typed_numbers(formulas, sheet):
    assert not typed_numbers(formulas[sheet]), sheet


def test_plan_types_only_its_input_and_the_backtest_years(formulas):
    ws = formulas["Plan"]
    k1 = next(r for r in range(1, 400) if ws[f"B{r}"].value == "Live")
    allowed = {"C4"} | {f"{c}{r}" for r in range(k1 + 1, k1 + 1 + len(L.BACKTEST_YEARS)) for c in "BCDEFGH"}
    assert not typed_numbers(ws, allowed)


def test_scenarios_type_only_levers_axes_and_table_workings(formulas):
    ws = formulas["Scenarios"]
    allowed = {f"C{r}" for r in range(4, 10)} | {f"{c}28" for c in "CDEFGHI"} | {f"B{r}" for r in range(29, 36)}
    for ref in ("C22:G24", "C29:I35", "K39:L58"):
        c1, r1, c2, r2 = range_boundaries(ref)
        allowed |= {f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)}
    assert not typed_numbers(ws, allowed)


def test_inputs_are_styled_as_inputs(formulas):
    ws = formulas["Assumptions"]
    for addr in ("C10", "C11", "C12", "C13", "C15", "C19"):
        assert ws[addr].fill.fgColor.rgb.endswith("FFF5D6"), addr


# --------------------------------------------------------- deliverables
def test_board_pack_pdf():
    pdf = ROOT / "docs" / "board_pack.pdf"
    assert pdf.exists() and len(re.findall(rb"/Type\s*/Page(?!s)", pdf.read_bytes())) >= 10


def test_workbook_opens_on_the_cover(values):
    assert values.active.title == "Cover"


def test_attribution_and_disclaimer_on_the_cover(values):
    text = " ".join(str(c.value) for row in values["Cover"].iter_rows() for c in row if c.value)
    assert "Not affiliated with or endorsed by Portland General Electric" in text
    assert "PUDL" in text and "CC BY 4.0" in text
