"""The FY2027 budget cycle workbook, held to its pandas twin.

budget_cycle/build_workbook.py writes every calculation as a formula and has
LibreOffice recalculate the file; budget_cycle/budget_model.py computes the same
cycle in pandas from the same inputs. These tests read the workbook's saved
values and compare them figure by figure: the consolidation of V0, V1 and V2,
the challenge-round bridge, every month of the three statements, and the 3+9
reforecast with its bridge. They also check the workbook's own controls, that
no calculation sheet holds a typed number, and that the extract the cycle starts
from still matches the Kestrel Bay workbook it was taken from.
"""
from __future__ import annotations

import json
import sys
import warnings

import openpyxl
import pandas as pd
import pytest

from workbook_io import ROOT

CYCLE = ROOT / "budget_cycle"
sys.path.insert(0, str(CYCLE))
import budget_model as reference  # noqa: E402  (not "reference": model/reference.py owns that name)

WORKBOOK = CYCLE / "Kestrel_Bay_FY2027_Budget.xlsx"
CELLS = json.loads((CYCLE / "cell_map.json").read_text(encoding="utf-8"))
CENT = 0.01


def _load(data_only: bool):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(WORKBOOK, data_only=data_only)


@pytest.fixture(scope="module")
def wb():
    return _load(True)


@pytest.fixture(scope="module")
def wbf():
    return _load(False)


@pytest.fixture(scope="module")
def ref():
    return reference.run()


def cell(wb, key: str):
    sheet, ref_ = CELLS[key].split("!")
    return wb[sheet][ref_].value


def row(wb, key: str) -> list[float]:
    sheet, rng = CELLS[key].split("!")
    return [c.value for c in wb[sheet][rng][0]]


# ---------------------------------------------------------------- consolidation
CONSOLIDATION = [
    ("base_revenue", lambda c: c["base"]["revenue"]),
    ("base_ebitda", lambda c: c["base"]["ebitda"]),
    ("target_revenue", lambda c: c["target"]["revenue"]),
    ("target_ebitda", lambda c: c["target"]["ebitda"]),
    *[(f"{v}_{m}", (lambda v, m: lambda c: c[v][m])(v, m))
      for v in ("v1", "v2") for m in ("revenue", "cogs", "freight", "gross_margin", "opex", "ebitda")],
    ("v1_ebitda_gap", lambda c: c["v1"]["ebitda"] - c["target"]["ebitda"]),
    ("v2_ebitda_gap", lambda c: c["v2"]["ebitda"] - c["target"]["ebitda"]),
    ("v2_revenue_gap", lambda c: c["v2"]["revenue"] - c["target"]["revenue"]),
    ("bridge_price", lambda c: c["bridge_v1_v2"]["price"]),
    *[(f"bridge_opex_{k}", (lambda d: lambda c: c["bridge_v1_v2"][f"opex {d}"])(d))
      for k, d in enumerate(reference.A["departments"])],
]


@pytest.mark.parametrize(("key", "expected"), CONSOLIDATION, ids=[k for k, _ in CONSOLIDATION])
def test_consolidation_matches_the_reference(wb, ref, key, expected):
    assert cell(wb, key) == pytest.approx(expected(ref["consolidation"]), abs=CENT)


def test_the_challenge_round_bridge_closes(wb, ref):
    c = ref["consolidation"]
    assert sum(c["bridge_v1_v2"].values()) == pytest.approx(c["v2"]["ebitda"] - c["v1"]["ebitda"], abs=1e-6)
    assert cell(wb, "bridge_total") == pytest.approx(c["v2"]["ebitda"] - c["v1"]["ebitda"], abs=CENT)


def test_the_price_needed_to_close_the_gap(wb, ref):
    b, c = ref["build"], ref["consolidation"]
    eff = reference.A["v1_submissions"]["price_effective_month"]
    needed = (c["target"]["ebitda"] - c["v2"]["ebitda"]) / b.loc[b.month >= eff, "revenue_v2"].sum()
    assert cell(wb, "price_to_close") == pytest.approx(needed, abs=1e-9)


def test_the_story_the_numbers_tell(ref):
    """What the README says about this cycle, held to the numbers."""
    c = ref["consolidation"]
    assert c["v1"]["ebitda"] < c["base"]["ebitda"], "bottom-up lands below last year: cost outruns price"
    assert c["v1"]["ebitda"] < c["v2"]["ebitda"] < c["target"]["ebitda"], "the challenge closes part of the gap"
    assert c["v2"]["revenue"] > c["target"]["revenue"], "revenue beats its target while EBITDA misses"


# ---------------------------------------------------------------- statements
STATEMENT_ROWS = ["revenue", "cogs", "freight", "gross_margin", "opex", "ebitda", "da", "ebit", "interest", "ebt",
                  "tax", "net_income", "capex", "cash", "ar", "inventory", "ppe", "ap", "revolver", "equity", "cfo",
                  "revolver_flow"]


@pytest.mark.parametrize("line", STATEMENT_ROWS)
def test_every_month_of_the_three_statements_matches(wb, ref, line):
    got = row(wb, f"stmt_{line}_row")
    want = ref["statements"][line].tolist()
    assert got == pytest.approx(want, abs=CENT), line


def test_the_balance_sheet_balances_every_month(wb, ref):
    assert max(abs(v) for v in row(wb, "stmt_balance_check_row")) < CENT
    assert ref["statements"].balance_check.abs().max() < 1e-6


def test_the_revolver_holds_minimum_cash(ref):
    s, bs = ref["statements"], reference.A["balance_sheet"]
    assert (s.revolver >= -1e-9).all()
    assert (s.cash >= bs["minimum_cash"] - 1e-6).all()
    assert ((s.revolver.abs() < 1e-6) | ((s.cash - bs["minimum_cash"]).abs() < 1e-6)).all(), \
        "cash above the minimum only once the revolver is repaid"


# ---------------------------------------------------------------- reforecast
def test_reforecast_ratios_match(wb, ref):
    rf = ref["reforecast"]
    for k, ch in enumerate(reference.A["channels"]):
        assert cell(wb, f"vol_ratio_{k}") == pytest.approx(rf["volume_ratio"][ch], abs=1e-9)
    for k, cat in enumerate(reference.A["categories"]):
        assert cell(wb, f"cost_ratio_{k}") == pytest.approx(rf["cost_ratio"][cat], abs=1e-9)
    for k, d in enumerate(reference.A["departments"]):
        assert cell(wb, f"opex_ratio_{k}") == pytest.approx(rf["opex_ratio"][d], abs=1e-9)


@pytest.mark.parametrize("line", ["revenue", "gross_margin", "opex", "ebitda"])
def test_reforecast_totals_match(wb, ref, line):
    assert cell(wb, f"rf_{line}") == pytest.approx(ref["reforecast"]["reforecast"][line], abs=CENT)


@pytest.mark.parametrize("step", ["q1", "volume", "cost", "opex"])
def test_reforecast_bridge_matches(wb, ref, step):
    names = {"q1": "q1_actual_vs_budget", "volume": "volume", "cost": "unit_cost", "opex": "opex"}
    assert cell(wb, f"bridge_rf_{step}") == pytest.approx(ref["reforecast"]["bridge"][names[step]], abs=CENT)


def test_the_reforecast_bridge_closes(ref):
    rf = ref["reforecast"]
    assert sum(rf["bridge"].values()) == pytest.approx(rf["reforecast"]["ebitda"] - rf["budget"]["ebitda"], abs=1e-6)


def test_q1_actuals_are_the_committed_extract(wb):
    q = pd.read_csv(CYCLE / "data" / "q1_sales_actuals.csv")
    ws = wb["Data_Q1"]
    got = [ws.cell(row=5 + i, column=5).value for i in range(len(q))]
    assert got == pytest.approx(q.revenue.tolist(), abs=1e-9)


# ---------------------------------------------------------------- the workbook itself
def test_every_check_in_the_workbook_passes(wb):
    sheet, rng = CELLS["checks_range"].split("!")
    statuses = [c[0].value for c in wb[sheet][rng]]
    assert statuses and set(statuses) == {"PASS"}
    assert cell(wb, "checks_passing") == int(CELLS["checks_total"]) == len(statuses)


def test_no_cell_holds_an_error(wb):
    errors = [f"{ws.title}!{c.coordinate}" for ws in wb.worksheets for r in ws.iter_rows() for c in r
              if isinstance(c.value, str) and c.value.startswith("#")]
    assert errors == []


CALC_SHEETS = ["Summary", "Consolidation", "Statements", "Reforecast", "Build", "Opex"]


@pytest.mark.parametrize("sheet", CALC_SHEETS)
def test_no_calculation_sheet_holds_a_typed_number(wbf, sheet):
    typed = [c.coordinate for r in wbf[sheet].iter_rows() for c in r
             if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
    assert typed == [], f"{sheet} has typed numbers at {typed[:5]}"


def test_the_extract_still_matches_the_kestrel_bay_workbook(tmp_path):
    """extract_base.py reads the tagged Kestrel Bay workbook; rerun it and
    require the committed CSVs back, byte for byte."""
    import importlib.util
    import shutil

    spec = importlib.util.spec_from_file_location("extract_base", CYCLE / "extract_base.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    committed = {p.name: p.read_bytes() for p in (CYCLE / "data").glob("*.csv") if not p.name.startswith("q1_")}
    mod.OUT = tmp_path
    mod.main()
    for name, data in committed.items():
        assert (tmp_path / name).read_bytes() == data, name
    shutil.rmtree(tmp_path, ignore_errors=True)


def test_q1_actuals_regenerate_identically(tmp_path):
    i = reference.inputs()
    sales, ox = reference.q1_actuals(i, reference.build(i))
    committed = pd.read_csv(CYCLE / "data" / "q1_sales_actuals.csv")
    pd.testing.assert_frame_equal(sales.reset_index(drop=True), committed, check_dtype=False)
    pd.testing.assert_frame_equal(ox.reset_index(drop=True), pd.read_csv(CYCLE / "data" / "q1_opex_actuals.csv"),
                                  check_dtype=False)


def test_the_readme_states_the_numbers_the_model_computes(ref):
    text = (CYCLE / "README.md").read_text(encoding="utf-8")
    c, rf, b = ref["consolidation"], ref["reforecast"], ref["build"]

    def m(x: float) -> str:
        return f"${x / 1e6:.2f}M"

    def k(x: float) -> str:
        return f"{'−' if x < 0 else ''}${abs(x) / 1e3:.0f}K"

    for v in ("v1", "v2"):
        assert m(c[v]["ebitda"]) in text and m(c[v]["revenue"]) in text
    assert m(rf["reforecast"]["ebitda"]) in text
    gap = c["target"]["ebitda"] - c["v2"]["ebitda"]
    eff = reference.A["v1_submissions"]["price_effective_month"]
    assert k(gap) in text
    assert f"{gap / b.loc[b.month >= eff, 'revenue_v2'].sum():.2%}" in text
    assert k(c["v2"]["ebitda"] - c["v1"]["ebitda"]) in text
    for step in rf["bridge"].values():
        assert k(step) in text
    assert k(rf["reforecast"]["ebitda"] - rf["budget"]["ebitda"]) in text
