"""
Build the lululemon DCF workbook.

    python valuation/build_workbook.py

Writes valuation/Lululemon_DCF_FY2025.xlsx. The filed figures (with the
accession number of each), the market inputs (with their sources) and the
scenario drivers are written to the input sheets; every other cell is a formula,
including the WACC-by-growth sensitivity grid. LibreOffice (headless) then opens
and saves the file so it carries calculated values, and tests/test_dcf.py holds
them to valuation/dcf_model.py.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OUT = HERE / "Lululemon_DCF_FY2025.xlsx"
FACTS = json.loads((HERE / "edgar_facts.json").read_text(encoding="utf-8"))
MARKET = json.loads((HERE / "market_inputs.json").read_text(encoding="utf-8"))
A = json.loads((HERE / "assumptions.json").read_text(encoding="utf-8"))
YEARS = list(FACTS["fiscal_years"])
FYEARS = A["explicit_years"]
SCEN = ["bear", "base", "bull"]
SCOL = {"bear": "C", "base": "D", "bull": "E"}

NAVY = "1F3A5F"
INPUT_FONT = Font(color="1F4E9E", bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
NOTE_FONT = Font(italic=True, color="555555")
BOLD = Font(bold=True)
MUSD = '#,##0.0,,;[Red](#,##0.0,,);"–"'
USD = '"$"#,##0.00'
PCT = "0.0%;[Red]-0.0%"
PCT2 = "0.00%;[Red]-0.00%"
X1 = '0.0"x"'
NUM2 = "0.00"

CELLS: dict[str, str] = {}


def title(ws, text, sub=None):
    ws["A1"], ws["A1"].font = text, TITLE_FONT
    if sub:
        ws["A2"], ws["A2"].font = sub, NOTE_FONT


def header(ws, row, labels, col=1):
    for j, label in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=label)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def inp(ws, ref, value, fmt=None, key=None):
    ws[ref] = value
    ws[ref].font, ws[ref].fill = INPUT_FONT, INPUT_FILL
    if fmt:
        ws[ref].number_format = fmt
    if key:
        CELLS[key] = f"{ws.title}!{ref}"


def put(ws, ref, formula, fmt=None, key=None, bold=False):
    ws[ref] = formula
    if fmt:
        ws[ref].number_format = fmt
    if bold:
        ws[ref].font = BOLD
    if key:
        CELLS[key] = f"{ws.title}!{ref}"


def widths(ws, w):
    for col, width in w.items():
        ws.column_dimensions[col].width = width


# ------------------------------------------------------------------ inputs
FACT_ROWS: dict[str, int] = {}


def data_sheet(wb):
    ws = wb.create_sheet("Data_EDGAR")
    title(ws, "lululemon athletica: filed figures (USD, shares) from SEC EDGAR XBRL",
          f"CIK {FACTS['cik']}. {FACTS['rule']}")
    header(ws, 4, ["Item", "XBRL tag", *YEARS, *[f"{y} accession" for y in YEARS]])
    r = 5
    for key, f in FACTS["facts"].items():
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=f["tag"])
        for j, y in enumerate(YEARS):
            v = f["values"][y]
            inp(ws, f"{get_column_letter(3 + j)}{r}", v["value"], "#,##0")
            ws.cell(row=r, column=8 + j, value=v["accession"])
        FACT_ROWS[key] = r
        r += 1
    widths(ws, {"A": 30, "B": 62, **{get_column_letter(c): 15 for c in range(3, 8)},
                **{get_column_letter(c): 22 for c in range(8, 13)}})


def F(key, year):
    """Reference to a filed figure."""
    return f"Data_EDGAR!${get_column_letter(3 + YEARS.index(year))}${FACT_ROWS[key]}"


INP: dict[str, str] = {}


def inputs_sheet(wb):
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs: market data and the three scenarios",
          "Blue on yellow is an input. Every market input names its source; every driver is a judgement, stated.")
    header(ws, 4, ["Market input", "Value", "Source"])
    rows = [
        ("valuation_date", "Valuation date", date.fromisoformat(MARKET["valuation_date"]), "yyyy-mm-dd", MARKET["why_that_date"]),
        ("first_mid", "Middle of FY2026", date.fromisoformat(A["first_forecast_midpoint"]), "yyyy-mm-dd",
         "cash flows arrive, on average, mid-year"),
        ("rf", "Risk-free rate", MARKET["risk_free_rate"]["value"], PCT2, MARKET["risk_free_rate"]["source"]),
        ("erp", "Equity risk premium", MARKET["equity_risk_premium"]["value"], PCT2, MARKET["equity_risk_premium"]["source"]),
        ("beta", "Beta", MARKET["beta"]["value"], NUM2, MARKET["beta"]["what"] + "; " + MARKET["beta"]["source"]),
        ("beta_alt", "Beta, alternative", MARKET["beta"]["alternative"]["value"], NUM2, MARKET["beta"]["alternative"]["what"]),
        ("debt", "Debt", MARKET["debt"]["value"], "#,##0", MARKET["debt"]["what"]),
        ("shares_common", "Common shares", MARKET["shares"]["common"], "#,##0", MARKET["shares"]["source"]),
        ("shares_exch", "Exchangeable shares", MARKET["shares"]["exchangeable"], "#,##0", MARKET["shares"]["what"]),
    ]
    for i, (key, label, value, fmt, source) in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=1, value=label)
        inp(ws, f"B{r}", value, fmt, key=f"input_{key}")
        ws.cell(row=r, column=3, value=source)
        INP[key] = f"Inputs!$B${r}"

    r0 = 16
    header(ws, r0, ["Driver", "Year", "Bear", "Base", "Bull"])
    r = r0 + 1
    for driver, label, fmt in [("revenue_growth", "Revenue growth", PCT), ("ebit_margin", "Operating (EBIT) margin", PCT),
                               ("capex_pct", "Capex, % of revenue", PCT2)]:
        for k, y in enumerate(FYEARS):
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=y)
            for s in SCEN:
                inp(ws, f"{SCOL[s]}{r}", A["scenarios"][s][driver][k], fmt)
            INP[f"{driver}_{k}"] = r
            r += 1
    for driver, label, fmt in [("da_pct", "D&A, % of revenue", PCT2), ("nwc_pct", "Operating working capital, % of revenue", PCT),
                               ("tax_rate", "Tax rate", PCT), ("terminal_growth", "Terminal growth", PCT2),
                               ("terminal_return_on_new_capital", "Return on new capital, terminal", PCT)]:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value="all years")
        for s in SCEN:
            inp(ws, f"{SCOL[s]}{r}", A["scenarios"][s][driver], fmt)
        INP[driver] = r
        r += 1
    ws.cell(row=r + 1, column=1, value="The story behind each case").font = BOLD
    for i, s in enumerate(SCEN):
        ws.cell(row=r + 2 + i, column=1, value=s.title())
        ws.cell(row=r + 2 + i, column=2, value=A["scenarios"][s]["story"])

    rs = r + 7
    ws.cell(row=rs, column=1, value="Sensitivity axes").font = BOLD
    ws.cell(row=rs + 1, column=1, value="WACC")
    for j, w in enumerate(A["sensitivity"]["wacc"]):
        inp(ws, f"{get_column_letter(2 + j)}{rs + 1}", w, PCT2)
    ws.cell(row=rs + 2, column=1, value="Terminal growth")
    for j, g in enumerate(A["sensitivity"]["terminal_growth"]):
        inp(ws, f"{get_column_letter(2 + j)}{rs + 2}", g, PCT2)
    INP["sens_wacc_row"], INP["sens_g_row"] = rs + 1, rs + 2
    widths(ws, {"A": 40, "B": 16, "C": 16, "D": 16, "E": 16})
    ws.column_dimensions["C"].width = 16


def drv(s, key, k=None):
    row = INP[f"{key}_{k}"] if k is not None else INP[key]
    return f"Inputs!${SCOL[s]}${row}"


# ------------------------------------------------------------------ history
HIST: dict[str, int] = {}


def history_sheet(wb):
    ws = wb.create_sheet("History")
    title(ws, "History: what the filings say, FY2021 to FY2025",
          "Ratios over the filed figures. 'Other operating items' is gross profit less SG&A, impairment and "
          "restructuring, and operating income: amortization and acquisition costs the filings report on other lines.")
    header(ws, 4, ["", *YEARS])
    lines = [
        ("revenue", "Revenue", lambda y, c: f"={F('revenue', y)}", MUSD),
        ("growth", "Revenue growth", lambda y, c: None if y == YEARS[0] else
         f"={c}5/{get_column_letter(ord(c) - 64 - 1)}5-1", PCT),
        ("gm", "Gross margin", lambda y, c: f"={F('gross_profit', y)}/{c}5", PCT),
        ("ebit", "Operating income (EBIT)", lambda y, c: f"={F('operating_income', y)}", MUSD),
        ("ebit_margin", "EBIT margin", lambda y, c: f"={c}8/{c}5", PCT),
        ("other_ops", "Other operating items", lambda y, c: f"={F('gross_profit', y)}-{F('sga', y)}-"
         f"{F('impairment_and_restructuring', y)}-{F('operating_income', y)}", MUSD),
        ("tax_rate", "Effective tax rate", lambda y, c: f"={F('income_tax', y)}/{F('pretax_income', y)}", PCT),
        ("da_pct", "D&A, % of revenue", lambda y, c: f"={F('depreciation_and_amortization', y)}/{c}5", PCT2),
        ("capex_pct", "Capex, % of revenue", lambda y, c: f"={F('capital_expenditure', y)}/{c}5", PCT2),
        ("nwc", "Operating working capital", lambda y, c: f"={F('receivables', y)}+{F('inventory', y)}+"
         f"{F('prepaid_and_other_current', y)}-{F('accounts_payable', y)}-{F('accrued_liabilities', y)}-"
         f"{F('income_taxes_payable', y)}", MUSD),
        ("nwc_pct", "Operating working capital, % of revenue", lambda y, c: f"={c}14/{c}5", PCT),
        ("sbc_pct", "Share-based compensation, % of revenue", lambda y, c: f"={F('share_based_compensation', y)}/{c}5", PCT2),
        ("fcf", "Cash from operations less capex", lambda y, c: f"={F('cash_from_operations', y)}-{F('capital_expenditure', y)}", MUSD),
        ("buybacks", "Share repurchases", lambda y, c: f"={F('share_repurchases', y)}", MUSD),
    ]
    for i, (key, label, fn, fmt) in enumerate(lines):
        r = 5 + i
        HIST[key] = r
        ws.cell(row=r, column=1, value=label)
        for j, y in enumerate(YEARS):
            col = get_column_letter(2 + j)
            f = fn(y, col)
            if f:
                put(ws, f"{col}{r}", f, fmt, key=f"hist_{key}_{y}")
    ws.cell(row=21, column=1, value="Money in USD millions.").font = NOTE_FONT
    widths(ws, {"A": 42, **{get_column_letter(c): 14 for c in range(2, 7)}})


# ------------------------------------------------------------------ DCF
DCF: dict[str, int] = {}


def dcf_sheet(wb, s):
    ws = wb.create_sheet(f"DCF_{s.title()}")
    title(ws, f"DCF, {s} case: free cash flow to the firm, FY2026 to FY2030",
          "Every cell is a formula over Inputs, History and Data_EDGAR. Money in USD millions.")
    header(ws, 4, ["", "FY2025 (filed)", *FYEARS])
    labels = ["Revenue", "Revenue growth", "EBIT", "EBIT margin", "Tax on EBIT", "NOPAT", "D&A", "Capex",
              "Operating working capital", "Increase in working capital", "Free cash flow to the firm",
              "Years from the valuation date", "Discount factor", "Present value"]
    keys = ["revenue", "growth", "ebit", "margin", "tax", "nopat", "da", "capex", "nwc", "dnwc", "fcff", "t", "df", "pv"]
    for i, (lab, key) in enumerate(zip(labels, keys, strict=True)):
        DCF[key] = 5 + i
        ws.cell(row=5 + i, column=1, value=lab)
    R = DCF
    put(ws, f"B{R['revenue']}", f"=History!F{HIST['revenue']}", MUSD)
    put(ws, f"B{R['ebit']}", f"=History!F{HIST['ebit']}", MUSD)
    put(ws, f"B{R['margin']}", f"=B{R['ebit']}/B{R['revenue']}", PCT)
    put(ws, f"B{R['nwc']}", f"=History!F{HIST['nwc']}", MUSD)
    for k, y in enumerate(FYEARS):
        c, p = get_column_letter(3 + k), get_column_letter(2 + k)
        put(ws, f"{c}{R['revenue']}", f"={p}{R['revenue']}*(1+{drv(s, 'revenue_growth', k)})", MUSD, key=f"{s}_revenue_{y}")
        put(ws, f"{c}{R['growth']}", f"={c}{R['revenue']}/{p}{R['revenue']}-1", PCT)
        put(ws, f"{c}{R['ebit']}", f"={c}{R['revenue']}*{drv(s, 'ebit_margin', k)}", MUSD, key=f"{s}_ebit_{y}")
        put(ws, f"{c}{R['margin']}", f"={c}{R['ebit']}/{c}{R['revenue']}", PCT)
        put(ws, f"{c}{R['tax']}", f"=-{c}{R['ebit']}*{drv(s, 'tax_rate')}", MUSD)
        put(ws, f"{c}{R['nopat']}", f"={c}{R['ebit']}+{c}{R['tax']}", MUSD, key=f"{s}_nopat_{y}")
        put(ws, f"{c}{R['da']}", f"={c}{R['revenue']}*{drv(s, 'da_pct')}", MUSD)
        put(ws, f"{c}{R['capex']}", f"=-{c}{R['revenue']}*{drv(s, 'capex_pct', k)}", MUSD)
        put(ws, f"{c}{R['nwc']}", f"={c}{R['revenue']}*{drv(s, 'nwc_pct')}", MUSD)
        put(ws, f"{c}{R['dnwc']}", f"=-({c}{R['nwc']}-{p}{R['nwc']})", MUSD)
        put(ws, f"{c}{R['fcff']}", f"={c}{R['nopat']}+{c}{R['da']}+{c}{R['capex']}+{c}{R['dnwc']}", MUSD,
            key=f"{s}_fcff_{y}", bold=True)
        put(ws, f"{c}{R['t']}", f"=({INP['first_mid']}-{INP['valuation_date']})/365.25+{k}", "0.000", key=f"t_{y}" if s == "base" else None)
        put(ws, f"{c}{R['df']}", f"=1/(1+$B$24)^{c}{R['t']}", "0.0000")
        put(ws, f"{c}{R['pv']}", f"={c}{R['fcff']}*{c}{R['df']}", MUSD)

    # valuation block
    rows = [
        ("Cost of equity = risk-free + beta x ERP (no debt, so the WACC)", f"={INP['rf']}+{INP['beta']}*{INP['erp']}", PCT2, "wacc"),
        ("Terminal growth", f"={drv(s, 'terminal_growth')}", PCT2, "g"),
        ("Return on new capital, terminal", f"={drv(s, 'terminal_return_on_new_capital')}", PCT, "ronic"),
        ("NOPAT, FY2031", f"=G{R['ebit']}*(1+B25)*(1-{drv(s, 'tax_rate')})", MUSD, "nopat_next"),
        ("Terminal value at end FY2030 = NOPAT x (1 - g/RONIC) / (WACC - g)", "=B27*(1-B25/B26)/(B24-B25)", MUSD, "terminal_value"),
        ("Present value of the terminal value", f"=B28/(1+B24)^(G{R['t']}+0.5)", MUSD, "pv_terminal"),
        ("Present value of FY2026-FY2030 cash flows", f"=SUM(C{R['pv']}:G{R['pv']})", MUSD, "pv_fcff"),
        ("Enterprise value", "=B29+B30", MUSD, "enterprise_value"),
        ("Cash (incl. restricted), FY2025 year end", f"={F('cash', YEARS[-1])}", MUSD, "cash"),
        ("Debt", f"=-{INP['debt']}", MUSD, "debt"),
        ("Equity value", "=B31+B32+B33", MUSD, "equity_value"),
        ("Shares: common plus exchangeable", f"={INP['shares_common']}+{INP['shares_exch']}", "#,##0", "shares"),
        ("Value per share (USD)", "=B34/B35", USD, "value_per_share"),
        ("Terminal value, share of enterprise value", "=B29/B31", PCT, "terminal_share_of_ev"),
        ("Implied exit multiple: terminal value / FY2030 EBITDA", f"=B28/(G{R['ebit']}+G{R['da']})", X1, "implied_terminal_ev_ebitda"),
        ("Enterprise value / FY2025 EBIT", f"=B31/{F('operating_income', YEARS[-1])}", X1, "ev_to_fy2025_ebit"),
    ]
    for i, (label, formula, fmt, key) in enumerate(rows):
        r = 24 + i
        ws.cell(row=r, column=1, value=label)
        put(ws, f"B{r}", formula, fmt, key=f"{s}_{key}", bold=key in ("enterprise_value", "equity_value", "value_per_share"))
    widths(ws, {"A": 62, "B": 16, **{get_column_letter(c): 14 for c in range(3, 8)}})


def sensitivity_sheet(wb):
    ws = wb.create_sheet("Sensitivity")
    title(ws, "Sensitivity, base case: value per share by WACC and terminal growth",
          "Each cell re-discounts the base case's free cash flows and terminal value at its own WACC and growth.")
    ws["A4"] = "WACC \\ terminal growth"
    ws["A4"].font = BOLD
    wr, gr = INP["sens_wacc_row"], INP["sens_g_row"]
    R = DCF
    for j in range(len(A["sensitivity"]["terminal_growth"])):
        col = get_column_letter(2 + j)
        put(ws, f"{col}4", f"=Inputs!{get_column_letter(2 + j)}${gr}", PCT2)
        ws[f"{col}4"].font, ws[f"{col}4"].fill = HEAD_FONT, HEAD_FILL
    for i in range(len(A["sensitivity"]["wacc"])):
        r = 5 + i
        put(ws, f"A{r}", f"=Inputs!{get_column_letter(2 + i)}${wr}", PCT2)
        ws[f"A{r}"].font = BOLD
        for j in range(len(A["sensitivity"]["terminal_growth"])):
            col = get_column_letter(2 + j)
            w, g = f"$A{r}", f"{col}$4"
            formula = (f"=(SUMPRODUCT(DCF_Base!$C${R['fcff']}:$G${R['fcff']},(1+{w})^(-DCF_Base!$C${R['t']}:$G${R['t']}))"
                       f"+DCF_Base!$G${R['ebit']}*(1+{g})*(1-{drv('base', 'tax_rate')})*(1-{g}/{drv('base', 'terminal_return_on_new_capital')})"
                       f"/({w}-{g})/(1+{w})^(DCF_Base!$G${R['t']}+0.5)"
                       f"+DCF_Base!$B$32+DCF_Base!$B$33)/DCF_Base!$B$35")
            put(ws, f"{col}{r}", formula, USD, key=f"sens_{i}_{j}")
    ws["A15"] = "Beta of Retail (Special Lines) instead of Apparel, base case"
    put(ws, "B15", f"=(SUMPRODUCT(DCF_Base!$C${R['fcff']}:$G${R['fcff']},(1+B16)^(-DCF_Base!$C${R['t']}:$G${R['t']}))"
                   f"+DCF_Base!$B$27*(1-DCF_Base!$B$25/DCF_Base!$B$26)/(B16-DCF_Base!$B$25)/(1+B16)^(DCF_Base!$G${R['t']}+0.5)"
                   f"+DCF_Base!$B$32+DCF_Base!$B$33)/DCF_Base!$B$35", USD, key="beta_alt_value_per_share")
    ws["A16"] = "  its cost of equity"
    put(ws, "B16", f"={INP['rf']}+{INP['beta_alt']}*{INP['erp']}", PCT2, key="beta_alt_wacc")
    widths(ws, {"A": 58, **{get_column_letter(c): 13 for c in range(2, 8)}})


def checks_sheet(wb):
    ws = wb.create_sheet("Checks")
    title(ws, "Checks", "Each row is PASS or FAIL, computed by formula.")
    header(ws, 4, ["Check", "Status"])
    checks = []
    for j, y in enumerate(YEARS):
        checks.append((f"{y}: gross profit = revenue - cost of goods sold, as filed",
                       f'=IF(ABS({F("gross_profit", y)}-({F("revenue", y)}-{F("cost_of_goods_sold", y)}))<1,"PASS","FAIL")'))
        checks.append((f"{y}: net income = pre-tax income - income tax, as filed",
                       f'=IF(ABS({F("net_income", y)}-({F("pretax_income", y)}-{F("income_tax", y)}))<1,"PASS","FAIL")'))
    for s in SCEN:
        sh = f"DCF_{s.title()}"
        checks.append((f"{s}: enterprise value = PV of cash flows + PV of terminal value",
                       f'=IF(ABS({sh}!B31-SUM({sh}!C18:G18)-{sh}!B29)<0.01,"PASS","FAIL")'))
        checks.append((f"{s}: WACC above terminal growth", f'=IF({sh}!B24>{sh}!B25,"PASS","FAIL")'))
    checks.append(("Bear < base < bull, per share", '=IF(AND(DCF_Bear!B36<DCF_Base!B36,DCF_Base!B36<DCF_Bull!B36),"PASS","FAIL")'))
    gi = A["sensitivity"]["terminal_growth"].index(A["scenarios"]["base"]["terminal_growth"])
    checks.append(("Sensitivity: lower WACC is always worth more",
                   "=IF(AND(" + ",".join(f"Sensitivity!{get_column_letter(2 + gi)}{5 + i}>Sensitivity!{get_column_letter(2 + gi)}{6 + i}"
                                        for i in range(len(A["sensitivity"]["wacc"]) - 1)) + '),"PASS","FAIL")'))
    for i, (label, formula) in enumerate(checks):
        r = 5 + i
        ws.cell(row=r, column=1, value=label)
        put(ws, f"B{r}", formula)
    last = 5 + len(checks) - 1
    CELLS["checks_range"] = f"Checks!B5:B{last}"
    CELLS["checks_total"] = str(len(checks))
    ws.cell(row=last + 2, column=1, value="Checks passing").font = BOLD
    put(ws, f"B{last + 2}", f'=COUNTIF(B5:B{last},"PASS")', "0", key="checks_passing", bold=True)
    widths(ws, {"A": 70, "B": 12})


def summary_sheet(wb):
    ws = wb.active
    ws.title = "Summary"
    title(ws, "lululemon athletica: what the business is worth, on its own filings",
          f"A discounted cash flow valuation as of {MARKET['valuation_date']}, from the FY2025 10-K. "
          "Not investment advice: a worked valuation whose every input is sourced and whose every output is a formula.")
    header(ws, 4, ["", "Bear", "Base", "Bull"])
    rows = [("Value per share (USD)", 36, USD), ("Enterprise value (USD M)", 31, MUSD), ("Equity value (USD M)", 34, MUSD),
            ("Terminal value, share of EV", 37, PCT), ("Implied exit EV / EBITDA", 38, X1),
            ("Cost of equity = WACC", 24, PCT2), ("Terminal growth", 25, PCT2)]
    for i, (label, src, fmt) in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=1, value=label)
        for j, s in enumerate(SCEN):
            put(ws, f"{get_column_letter(2 + j)}{r}", f"=DCF_{s.title()}!B{src}", fmt, bold=i == 0)
    ws["A14"] = "FY2025 as filed"
    ws["A14"].font = BOLD
    for i, (label, ref, fmt) in enumerate([("Revenue (USD M)", f"History!F{HIST['revenue']}", MUSD),
                                           ("Revenue growth", f"History!F{HIST['growth']}", PCT),
                                           ("EBIT margin", f"History!F{HIST['ebit_margin']}", PCT),
                                           ("Operating working capital, % of revenue", f"History!F{HIST['nwc_pct']}", PCT)]):
        ws.cell(row=15 + i, column=1, value=label)
        put(ws, f"B{15 + i}", f"={ref}", fmt)
    ws["A20"] = "Checks passing"
    put(ws, "B20", "=Checks!B" + str(5 + int(CELLS["checks_total"]) + 1), "0")
    ws["C20"] = f"of {CELLS['checks_total']}"
    widths(ws, {"A": 44, "B": 16, "C": 16, "D": 16})


def recalculate(path: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "out"
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                       check=True, capture_output=True, timeout=300)
        shutil.copyfile(out / path.name, path)


def main() -> int:
    wb = Workbook()
    wb.active.title = "Summary"
    data_sheet(wb)
    inputs_sheet(wb)
    history_sheet(wb)
    for s in SCEN:
        dcf_sheet(wb, s)
    sensitivity_sheet(wb)
    checks_sheet(wb)
    summary_sheet(wb)
    order = ["Summary", "Inputs", "DCF_Base", "DCF_Bear", "DCF_Bull", "Sensitivity", "History", "Data_EDGAR", "Checks"]
    wb._sheets = [wb[name] for name in order]
    wb.save(OUT)
    recalculate(OUT)
    (HERE / "cell_map.json").write_text(json.dumps(CELLS, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {OUT.name}: {len(CELLS)} cells mapped")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
