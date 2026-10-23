"""
Build the Kestrel Bay FY2027 budget-cycle workbook.

    python budget_cycle/build_workbook.py

Writes budget_cycle/Kestrel_Bay_FY2027_Budget.xlsx with every calculation as an
Excel formula over the input and data sheets, then has LibreOffice (headless)
open and save it so the file carries calculated values. Nothing is pasted in:
the inputs are assumptions.json and data/*.csv, written to the Assumptions,
Data_Base and Data_Q1 sheets; everything else is a formula.

tests/test_budget_cycle.py reads the saved values and holds them to
budget_cycle/budget_model.py, which computes the same cycle in pandas.
Uses only functions Excel and LibreOffice share (SUMIFS, SUMPRODUCT, INDEX,
MATCH, EOMONTH), so it recalculates the same in either.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
OUT = HERE / "Kestrel_Bay_FY2027_Budget.xlsx"
A = json.loads((HERE / "assumptions.json").read_text(encoding="utf-8"))

NAVY = "1F3A5F"
INPUT_FONT = Font(color="1F4E9E", bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
SECTION_FONT = Font(bold=True, size=11, color=NAVY)
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="9AA5B1")
MONEY = '#,##0;[Red](#,##0);"–"'
MONEY2 = '#,##0.00;[Red](#,##0.00);"–"'
PCT = "0.0%;[Red]-0.0%"
PCT2 = "0.00%;[Red]-0.00%"
NUM1 = "#,##0.0"

N_SEG = 24
B0, B1 = 5, 5 + N_SEG * 12 - 1  # Build sheet rows
Q0, Q1R = 5, 5 + N_SEG * 3 - 1  # Data_Q1 sales rows
QO0, QO1 = 80, 91  # Data_Q1 opex rows

CELLS: dict[str, str] = {}  # published figure -> 'Sheet!A1', for the tests


def title(ws, text: str, sub: str | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="555555")


def header(ws, row: int, labels: list[str], col: int = 1) -> None:
    for j, label in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=label)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def inp(cell, value, fmt=None) -> None:
    cell.value = value
    cell.font, cell.fill = INPUT_FONT, INPUT_FILL
    if fmt:
        cell.number_format = fmt


def put(ws, ref: str, formula: str, fmt: str | None = None, key: str | None = None, bold: bool = False) -> None:
    ws[ref] = formula
    if fmt:
        ws[ref].number_format = fmt
    if bold:
        ws[ref].font = TOTAL_FONT
    if key:
        CELLS[key] = f"{ws.title}!{ref}"


def widths(ws, w: dict[str, float]) -> None:
    for col, width in w.items():
        ws.column_dimensions[col].width = width


# ------------------------------------------------------------------ inputs
def assumptions(wb) -> None:
    ws = wb.create_sheet("Assumptions")
    title(ws, "Assumptions: every input to the FY2027 cycle",
          "Blue on yellow is an input. Every other number in the workbook is a formula over these and the data sheets.")
    v1, v2, bs = A["v1_submissions"], A["v2_challenge"], A["balance_sheet"]
    ws["B3"] = "Fiscal year"
    inp(ws["C3"], A["fiscal_year"], "0")
    ws["A4"] = "V0  Top-down targets"
    ws["A4"].font = SECTION_FONT
    ws["B5"], ws["B6"] = "Revenue growth on the last twelve months", "EBITDA margin"
    inp(ws["C5"], A["targets"]["revenue_growth"], PCT)
    inp(ws["C6"], A["targets"]["ebitda_margin"], PCT)

    ws["A8"] = "V1  Category submissions (sales and procurement); V2 challenge"
    ws["A8"].font = SECTION_FONT
    header(ws, 9, ["Category", "V1 price increase", "V1 cost inflation", "V2 extra price, foodservice"])
    for k, cat in enumerate(A["categories"]):
        r = 10 + k
        ws.cell(row=r, column=1, value=cat)
        inp(ws.cell(row=r, column=2), v1["price_increase"][cat], PCT)
        inp(ws.cell(row=r, column=3), v1["cost_inflation"][cat], PCT)
        inp(ws.cell(row=r, column=4), v2["foodservice_extra_price"].get(cat, 0.0), PCT)

    ws["A17"] = "V1  Channel submissions (sales)"
    ws["A17"].font = SECTION_FONT
    header(ws, 18, ["Channel", "V1 volume growth", "Foodservice (1 = yes)"])
    for k, ch in enumerate(A["channels"]):
        r = 19 + k
        ws.cell(row=r, column=1, value=ch)
        inp(ws.cell(row=r, column=2), v1["volume_growth"][ch], PCT)
        inp(ws.cell(row=r, column=3), 1 if ch in v2["foodservice_channels"] else 0, "0")

    ws["A24"] = "V1  Department submissions; V2 challenge"
    ws["A24"].font = SECTION_FONT
    header(ws, 25, ["Department", "V1 opex growth", "V2 growth cap", "V2 cut to base"])
    for k, d in enumerate(A["departments"]):
        r = 26 + k
        ws.cell(row=r, column=1, value=d)
        inp(ws.cell(row=r, column=2), v1["opex_growth"][d], PCT)
        cap = v2["opex_growth_cap"].get(d)
        inp(ws.cell(row=r, column=3), cap if cap is not None else None, PCT)
        inp(ws.cell(row=r, column=4), v2["opex_base_cut"].get(d), PCT)

    ws["A31"] = "V1  New hires; V2 timing"
    ws["A31"].font = SECTION_FONT
    header(ws, 32, ["Department", "FTE", "Annual loaded cost", "Start month V1", "Start month V2", "Start used in V2"])
    for k, h in enumerate(v1["hires"]):
        r = 33 + k
        ws.cell(row=r, column=1, value=h["dept"])
        inp(ws.cell(row=r, column=2), h["fte"], "0")
        inp(ws.cell(row=r, column=3), h["annual_cost"], MONEY)
        inp(ws.cell(row=r, column=4), h["start_month"], "0")
        inp(ws.cell(row=r, column=5), v2["hire_start_month"].get(h["dept"]), "0")
        put(ws, f"F{r}", f'=IF(E{r}="",D{r},E{r})', "0")

    ws["A36"] = "Timing and freight"
    ws["A36"].font = SECTION_FONT
    ws["B37"], ws["B38"] = "Price increases take effect in month", "Freight inflation"
    inp(ws["C37"], v1["price_effective_month"], "0")
    inp(ws["C38"], v1["freight_inflation"], PCT)

    ws["A40"] = "Balance sheet and financing"
    ws["A40"].font = SECTION_FONT
    labels = [("Opening cash", "opening_cash", MONEY), ("Minimum cash", "minimum_cash", MONEY),
              ("Opening revolver", "opening_revolver", MONEY), ("Revolver rate", "revolver_rate", PCT2),
              ("Opening PP&E, net", "opening_ppe", MONEY), ("D&A on existing assets, per month", "existing_da_per_month", MONEY),
              ("Useful life of new capex, months", "useful_life_months", "0"), ("Tax rate", "tax_rate", PCT),
              ("Days sales outstanding", "dso", "0"), ("Days inventory outstanding", "dio", "0"),
              ("Days payables outstanding", "dpo", "0")]
    for k, (label, key, fmt) in enumerate(labels):
        r = 41 + k
        ws.cell(row=r, column=2, value=label)
        inp(ws.cell(row=r, column=3), bs[key], fmt)
    ws["A53"] = "Capex by month"
    ws["A53"].font = SECTION_FONT
    for m in range(1, 13):
        c = ws.cell(row=54, column=1 + m, value=m)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        inp(ws.cell(row=55, column=1 + m), float(bs["capex"].get(str(m), 0)), MONEY)
    ws["A55"] = "Capex"
    widths(ws, {"A": 26, "B": 40, "C": 18, "D": 18, "E": 16, "F": 16})


def data_sheets(wb) -> None:
    base = pd.read_csv(DATA / "base_segments.csv")
    season = pd.read_csv(DATA / "seasonality.csv")
    opex = pd.read_csv(DATA / "opex_base.csv")
    wc = pd.read_csv(DATA / "opening_wc.csv")
    ws = wb.create_sheet("Data_Base")
    title(ws, "Data: the budget's starting point (synthetic Kestrel Bay)",
          "From budget_cycle/data/*.csv, extracted from the Kestrel Bay workbook by extract_base.py. Loaded unchanged.")
    ws["A3"] = "Last twelve months of actuals, Sep 2025 to Aug 2026, by category and channel"
    header(ws, 4, ["Category", "Channel", "Cases", "Net revenue", "COGS", "Freight"])
    for k, r in base.iterrows():
        row = 5 + k
        for j, v in enumerate([r.category, r.channel, r.cases, r.revenue, r.cogs, r.freight]):
            ws.cell(row=row, column=1 + j, value=v).number_format = MONEY if j >= 2 else "General"
    ws["A30"] = "Seasonality: each category's share of its calendar-2025 cases, by month"
    header(ws, 31, ["Category"] + [f"m{m:02d}" for m in range(1, 13)])
    for k, r in season.iterrows():
        row = 32 + k
        ws.cell(row=row, column=1, value=r.category)
        for m in range(1, 13):
            ws.cell(row=row, column=1 + m, value=float(r[f"m{m:02d}"])).number_format = "0.0000"
    ws["A39"] = "Operating expenses, last twelve months"
    header(ws, 40, ["Department", "LTM opex"])
    for k, r in opex.iterrows():
        ws.cell(row=41 + k, column=1, value=r.dept)
        ws.cell(row=41 + k, column=2, value=float(r.amount)).number_format = MONEY
    ws["A46"] = "Working capital at the latest month-end (Aug 2026), which opens the budget's balance sheet"
    header(ws, 47, ["Account", "Balance"])
    for k, r in wc.iterrows():
        ws.cell(row=48 + k, column=1, value=r.account)
        ws.cell(row=48 + k, column=2, value=float(r.balance)).number_format = MONEY
    widths(ws, {"A": 22, "B": 22, "C": 12, "D": 14, "E": 14, "F": 12})

    q = pd.read_csv(DATA / "q1_sales_actuals.csv")
    qo = pd.read_csv(DATA / "q1_opex_actuals.csv")
    ws = wb.create_sheet("Data_Q1")
    title(ws, "Data: Q1 FY2027 actuals (synthetic)",
          "From data/q1_*.csv: the approved budget moved by stated deviations and seeded noise (budget_model.py --q1).")
    header(ws, 4, ["Month", "Category", "Channel", "Cases", "Net revenue", "COGS", "Freight"])
    for k, r in q.iterrows():
        row = Q0 + k
        for j, v in enumerate([int(r.month), r.category, r.channel, r.cases, r.revenue, r.cogs, r.freight]):
            ws.cell(row=row, column=1 + j, value=v).number_format = MONEY2 if j >= 3 else "General"
    header(ws, QO0 - 1, ["Month", "Department", "Opex"])
    for k, r in qo.iterrows():
        row = QO0 + k
        for j, v in enumerate([int(r.month), r.dept, r.amount]):
            ws.cell(row=row, column=1 + j, value=v).number_format = MONEY2 if j == 2 else "General"
    widths(ws, {"A": 8, "B": 22, "C": 22, "D": 12, "E": 14, "F": 14, "G": 12})


# ------------------------------------------------------------- calculation
def build_sheet(wb) -> None:
    ws = wb.create_sheet("Build")
    title(ws, "Build: segment x month, every version",
          "24 category x channel segments x 12 months. V1 and V2 share volumes and costs; V2 adds the challenge-round price. "
          "Columns P:S are the 3+9 reforecast.")
    header(ws, 4, ["Month", "Segment", "Category", "Channel", "Season share", "Cases", "Base price / case",
                   "V1 price uplift", "V2 extra price", "Revenue V1", "Revenue V2", "Cost / case", "COGS",
                   "Freight / case", "Freight", "RF cases", "RF revenue", "RF COGS", "RF freight"])
    cats = "Assumptions!$A$10:$A$15"
    chans = "Assumptions!$A$19:$A$22"
    for r in range(B0, B1 + 1):
        f = {
            "A": f"=MOD(ROW()-{B0},12)+1",
            "B": f"=INT((ROW()-{B0})/12)+1",
            "C": f"=INDEX(Data_Base!$A$5:$A$28,B{r})",
            "D": f"=INDEX(Data_Base!$B$5:$B$28,B{r})",
            "E": f"=INDEX(Data_Base!$B$32:$M$37,MATCH(C{r},Data_Base!$A$32:$A$37,0),A{r})",
            "F": f"=INDEX(Data_Base!$C$5:$C$28,B{r})*(1+INDEX(Assumptions!$B$19:$B$22,MATCH(D{r},{chans},0)))*E{r}",
            "G": f"=INDEX(Data_Base!$D$5:$D$28,B{r})/INDEX(Data_Base!$C$5:$C$28,B{r})",
            "H": f"=IF(A{r}>=Assumptions!$C$37,INDEX(Assumptions!$B$10:$B$15,MATCH(C{r},{cats},0)),0)",
            "I": (f"=IF(AND(A{r}>=Assumptions!$C$37,INDEX(Assumptions!$C$19:$C$22,MATCH(D{r},{chans},0))=1),"
                  f"INDEX(Assumptions!$D$10:$D$15,MATCH(C{r},{cats},0)),0)"),
            "J": f"=F{r}*G{r}*(1+H{r})",
            "K": f"=J{r}*(1+I{r})",
            "L": (f"=INDEX(Data_Base!$E$5:$E$28,B{r})/INDEX(Data_Base!$C$5:$C$28,B{r})"
                  f"*(1+INDEX(Assumptions!$C$10:$C$15,MATCH(C{r},{cats},0)))"),
            "M": f"=F{r}*L{r}",
            "N": f"=INDEX(Data_Base!$F$5:$F$28,B{r})/INDEX(Data_Base!$C$5:$C$28,B{r})*(1+Assumptions!$C$38)",
            "O": f"=F{r}*N{r}",
            "P": (f"=IF(A{r}<=3,SUMIFS(Data_Q1!$D${Q0}:$D${Q1R},Data_Q1!$A${Q0}:$A${Q1R},A{r},"
                  f"Data_Q1!$B${Q0}:$B${Q1R},C{r},Data_Q1!$C${Q0}:$C${Q1R},D{r}),"
                  f"F{r}*INDEX(Reforecast!$D$6:$D$9,MATCH(D{r},Reforecast!$A$6:$A$9,0)))"),
            "Q": (f"=IF(A{r}<=3,SUMIFS(Data_Q1!$E${Q0}:$E${Q1R},Data_Q1!$A${Q0}:$A${Q1R},A{r},"
                  f"Data_Q1!$B${Q0}:$B${Q1R},C{r},Data_Q1!$C${Q0}:$C${Q1R},D{r}),P{r}*G{r}*(1+H{r})*(1+I{r}))"),
            "R": (f"=IF(A{r}<=3,SUMIFS(Data_Q1!$F${Q0}:$F${Q1R},Data_Q1!$A${Q0}:$A${Q1R},A{r},"
                  f"Data_Q1!$B${Q0}:$B${Q1R},C{r},Data_Q1!$C${Q0}:$C${Q1R},D{r}),"
                  f"P{r}*L{r}*INDEX(Reforecast!$D$12:$D$17,MATCH(C{r},Reforecast!$A$12:$A$17,0)))"),
            "S": (f"=IF(A{r}<=3,SUMIFS(Data_Q1!$G${Q0}:$G${Q1R},Data_Q1!$A${Q0}:$A${Q1R},A{r},"
                  f"Data_Q1!$B${Q0}:$B${Q1R},C{r},Data_Q1!$C${Q0}:$C${Q1R},D{r}),P{r}*N{r})"),
        }
        fmts = {"E": "0.0000", "F": NUM1, "G": MONEY2, "H": PCT, "I": PCT, "J": MONEY, "K": MONEY, "L": MONEY2,
                "M": MONEY, "N": MONEY2, "O": MONEY, "P": NUM1, "Q": MONEY, "R": MONEY, "S": MONEY}
        for col, formula in f.items():
            ws[f"{col}{r}"] = formula
            if col in fmts:
                ws[f"{col}{r}"].number_format = fmts[col]
    ws.freeze_panes = "E5"
    widths(ws, {"C": 16, "D": 20, **{get_column_letter(c): 13 for c in range(5, 20)}})


def opex_sheet(wb) -> None:
    ws = wb.create_sheet("Opex")
    title(ws, "Opex: department submissions by month",
          "Base = last twelve months / 12 x (1 + growth) [x (1 - cut) in V2], plus hires from their start month.")
    for m in range(1, 13):
        put(ws, f"{get_column_letter(2 + m)}5", "=COLUMN()-2", "0")
    ws["O5"] = "FY"
    for blk, (top, label) in enumerate([(6, "V1 bottom-up"), (13, "V2 approved")]):
        ws.cell(row=top - 1, column=1, value=label).font = SECTION_FONT
        ws.cell(row=top - 1, column=2, value="Base / month").font = TOTAL_FONT
        for k, d in enumerate(A["departments"]):
            r = top + k
            ws.cell(row=r, column=1, value=d)
            dep = f"MATCH($A{r},Assumptions!$A$26:$A$29,0)"
            ltm = f"INDEX(Data_Base!$B$41:$B$44,MATCH($A{r},Data_Base!$A$41:$A$44,0))"
            g = f"INDEX(Assumptions!$B$26:$B$29,{dep})"
            if blk == 0:
                put(ws, f"B{r}", f"={ltm}/12*(1+{g})", MONEY)
                start = "Assumptions!$D$33:$D$34"
            else:
                cap = f"INDEX(Assumptions!$C$26:$C$29,{dep})"
                cut = f"INDEX(Assumptions!$D$26:$D$29,{dep})"
                put(ws, f"B{r}", f'={ltm}/12*(1+IF({cap}="",{g},MIN({g},{cap})))*(1-N({cut}))', MONEY)
                start = "Assumptions!$F$33:$F$34"
            for m in range(1, 13):
                col = get_column_letter(2 + m)
                put(ws, f"{col}{r}",
                    f"=$B{r}+SUMPRODUCT((Assumptions!$A$33:$A$34=$A{r})*({start}<={col}$5)"
                    f"*Assumptions!$B$33:$B$34*Assumptions!$C$33:$C$34)/12", MONEY)
            put(ws, f"O{r}", f"=SUM(C{r}:N{r})", MONEY, bold=True)
        tr = top + 4
        ws.cell(row=tr, column=1, value=f"Total {label.split()[0]}").font = TOTAL_FONT
        for c in range(3, 16):
            col = get_column_letter(c)
            put(ws, f"{col}{tr}", f"=SUM({col}{top}:{col}{top + 3})", MONEY, bold=True)
    CELLS["opex_v1_total"] = "Opex!O10"
    CELLS["opex_v2_total"] = "Opex!O17"
    widths(ws, {"A": 24, "B": 14, **{get_column_letter(c): 12 for c in range(3, 16)}})


def consolidation(wb) -> None:
    ws = wb.create_sheet("Consolidation")
    title(ws, "Consolidation: top-down meets bottom-up",
          "V0 is what leadership asked for, V1 what the business submitted, V2 what the challenge round approved.")
    header(ws, 5, ["", "Last 12 months", "V0 target", "V1 bottom-up", "V2 approved"])
    rows = ["Net revenue", "COGS", "Freight", "Gross margin", "Opex", "EBITDA", "EBITDA margin"]
    for k, label in enumerate(rows):
        ws.cell(row=6 + k, column=1, value=label).font = TOTAL_FONT if label in ("Gross margin", "EBITDA") else Font()
    b = f"Build!${{c}}${B0}:${{c}}${B1}"
    put(ws, "B6", "=SUM(Data_Base!D5:D28)", MONEY, "base_revenue")
    put(ws, "C6", "=B6*(1+Assumptions!C5)", MONEY, "target_revenue")
    put(ws, "D6", f"=SUM({b.format(c='J')})", MONEY, "v1_revenue")
    put(ws, "E6", f"=SUM({b.format(c='K')})", MONEY, "v2_revenue")
    put(ws, "B7", "=SUM(Data_Base!E5:E28)", MONEY)
    put(ws, "B8", "=SUM(Data_Base!F5:F28)", MONEY)
    for col, v in (("D", "v1"), ("E", "v2")):
        put(ws, f"{col}7", f"=SUM({b.format(c='M')})", MONEY, f"{v}_cogs")
        put(ws, f"{col}8", f"=SUM({b.format(c='O')})", MONEY, f"{v}_freight")
    for col in "BDE":
        put(ws, f"{col}9", f"={col}6-{col}7-{col}8", MONEY, bold=True)
    put(ws, "B10", "=SUM(Data_Base!B41:B44)", MONEY)
    put(ws, "D10", "=Opex!O10", MONEY, "v1_opex")
    put(ws, "E10", "=Opex!O17", MONEY, "v2_opex")
    for col in "BDE":
        put(ws, f"{col}11", f"={col}9-{col}10", MONEY, bold=True)
    put(ws, "C11", "=C6*Assumptions!C6", MONEY, "target_ebitda")
    CELLS.update({"base_ebitda": "Consolidation!B11", "v1_ebitda": "Consolidation!D11",
                  "v2_ebitda": "Consolidation!E11", "v1_gross_margin": "Consolidation!D9",
                  "v2_gross_margin": "Consolidation!E9"})
    for col in "BCDE":
        put(ws, f"{col}12", f"={col}11/{col}6", PCT2)

    ws["A14"] = "Gap to target"
    ws["A14"].font = SECTION_FONT
    ws["A15"], ws["A16"] = "Revenue", "EBITDA"
    for col, v in (("D", "v1"), ("E", "v2")):
        put(ws, f"{col}15", f"={col}6-$C$6", MONEY, f"{v}_revenue_gap")
        put(ws, f"{col}16", f"={col}11-$C$11", MONEY, f"{v}_ebitda_gap")

    ws["A18"] = "Challenge round: V1 to V2 EBITDA"
    ws["A18"].font = SECTION_FONT
    ws["A19"] = "Foodservice price on beef and seafood"
    put(ws, "B19", "=E6-D6", MONEY, "bridge_price")
    for k, d in enumerate(A["departments"]):
        r = 20 + k
        ws.cell(row=r, column=1, value=f"Opex: {d}")
        put(ws, f"B{r}", f"=Opex!O{6 + k}-Opex!O{13 + k}", MONEY, f"bridge_opex_{k}")
    ws["A24"] = "V1 to V2, total"
    put(ws, "B24", "=SUM(B19:B23)", MONEY, "bridge_total", bold=True)

    ws["A26"] = "What closing the rest would take"
    ws["A26"].font = SECTION_FONT
    ws["A27"] = "Further price on every line from the effective month, to reach the EBITDA target"
    put(ws, "B27", f"=-E16/SUMIFS({b.format(c='K')},{b.format(c='A')},\">=\"&Assumptions!C37)", PCT2,
        "price_to_close")
    widths(ws, {"A": 62, "B": 16, "C": 16, "D": 16, "E": 16})


def statements(wb) -> None:
    ws = wb.create_sheet("Statements")
    title(ws, "Three statements: the approved budget (V2), monthly",
          "Interest accrues on the opening revolver, so there is no circular reference. "
          "The revolver draws to hold minimum cash and repays from any surplus.")
    ws["B4"] = "Opening"
    for m in range(1, 13):
        put(ws, f"{get_column_letter(2 + m)}4", "=COLUMN()-2", "0")
    ws["O4"] = "FY2027"
    for c in range(1, 16):
        ws.cell(row=4, column=c).font, ws.cell(row=4, column=c).fill = HEAD_FONT, HEAD_FILL
    labels = {5: "Income statement", 6: "Net revenue", 7: "COGS", 8: "Freight", 9: "Gross margin", 10: "Opex",
              11: "EBITDA", 12: "Depreciation & amortisation", 13: "EBIT", 14: "Interest", 15: "Pre-tax income",
              16: "Tax", 17: "Net income", 19: "Capex", 20: "Days in month",
              22: "Balance sheet", 23: "Cash", 24: "Accounts receivable", 25: "Inventory", 26: "PP&E, net",
              27: "Total assets", 28: "Accounts payable", 30: "Revolver", 31: "Equity",
              32: "Total liabilities and equity", 33: "Balance check (assets - L&E)",
              34: "Cash flow", 35: "Net income", 36: "D&A", 37: "(Increase) in receivables",
              38: "(Increase) in inventory", 39: "Increase in payables", 40: "Cash from operations",
              41: "Capex", 42: "Cash before financing", 43: "Revolver draw / (repayment)", 44: "Net change in cash",
              45: "Cash check (vs balance sheet)"}
    for r, label in labels.items():
        ws.cell(row=r, column=1, value=label)
        if r in (5, 22, 34):
            ws.cell(row=r, column=1).font = SECTION_FONT
        if r in (9, 11, 17, 27, 32, 40):
            ws.cell(row=r, column=1).font = TOTAL_FONT
    b = f"Build!${{c}}${B0}:${{c}}${B1}"
    wc = "INDEX(Data_Base!$B$48:$B$50,MATCH(\"{a}\",Data_Base!$A$48:$A$50,0))"
    put(ws, "B23", "=Assumptions!C41", MONEY)
    put(ws, "B24", "=" + wc.format(a="Accounts receivable"), MONEY)
    put(ws, "B25", "=" + wc.format(a="Inventory"), MONEY)
    put(ws, "B26", "=Assumptions!C45", MONEY)
    put(ws, "B27", "=SUM(B23:B26)", MONEY, bold=True)
    put(ws, "B28", "=" + wc.format(a="Accounts payable"), MONEY)
    put(ws, "B30", "=Assumptions!C43", MONEY)
    put(ws, "B31", "=B27-B28-B30", MONEY)
    put(ws, "B32", "=B28+B30+B31", MONEY, bold=True)
    put(ws, "B33", "=B27-B32", MONEY2)
    for m in range(1, 13):
        c = get_column_letter(2 + m)
        p = get_column_letter(1 + m)
        f = {
            6: f"=SUMIFS({b.format(c='K')},{b.format(c='A')},{c}$4)",
            7: f"=SUMIFS({b.format(c='M')},{b.format(c='A')},{c}$4)",
            8: f"=SUMIFS({b.format(c='O')},{b.format(c='A')},{c}$4)",
            9: f"={c}6-{c}7-{c}8",
            10: f"=INDEX(Opex!$C$17:$N$17,{c}$4)",
            11: f"={c}9-{c}10",
            12: f'=Assumptions!$C$46+SUMIF(Assumptions!$B$54:$M$54,"<"&{c}$4,Assumptions!$B$55:$M$55)/Assumptions!$C$47',
            13: f"={c}11-{c}12",
            14: f"={p}30*Assumptions!$C$44/12",
            15: f"={c}13-{c}14",
            16: f"={c}15*Assumptions!$C$48",
            17: f"={c}15-{c}16",
            19: f"=INDEX(Assumptions!$B$55:$M$55,{c}$4)",
            20: f"=DAY(EOMONTH(DATE(Assumptions!$C$3,{c}$4,1),0))",
            23: f"={c}42+{c}43",
            24: f"={c}6/{c}20*Assumptions!$C$49",
            25: f"={c}7/{c}20*Assumptions!$C$50",
            26: f"={p}26+{c}19-{c}12",
            27: f"=SUM({c}23:{c}26)",
            28: f"={c}7/{c}20*Assumptions!$C$51",
            30: f"={p}30+{c}43",
            31: f"={p}31+{c}17",
            32: f"={c}28+{c}30+{c}31",
            33: f"={c}27-{c}32",
            35: f"={c}17",
            36: f"={c}12",
            37: f"=-({c}24-{p}24)",
            38: f"=-({c}25-{p}25)",
            39: f"={c}28-{p}28",
            40: f"=SUM({c}35:{c}39)",
            41: f"=-{c}19",
            42: f"={p}23+{c}40+{c}41",
            43: f"=MAX(Assumptions!$C$42-{c}42,-{p}30)",
            44: f"={c}40+{c}41+{c}43",
            45: f"={c}23-{p}23-{c}44",
        }
        for r, formula in f.items():
            put(ws, f"{c}{r}", formula, "0" if r == 20 else (MONEY2 if r in (33, 45) else MONEY),
                bold=r in (9, 11, 17, 27, 32, 40))
    flows = [6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 19, 35, 36, 37, 38, 39, 40, 41, 43, 44]
    for r in flows:
        put(ws, f"O{r}", f"=SUM(C{r}:N{r})", MONEY, bold=True)
    for r in (23, 24, 25, 26, 27, 28, 30, 31, 32):
        put(ws, f"O{r}", f"=N{r}", MONEY, bold=True)
    keys = {6: "revenue", 7: "cogs", 8: "freight", 9: "gross_margin", 10: "opex", 11: "ebitda", 12: "da",
            13: "ebit", 14: "interest", 15: "ebt", 16: "tax", 17: "net_income", 19: "capex", 23: "cash",
            24: "ar", 25: "inventory", 26: "ppe", 28: "ap", 30: "revolver", 31: "equity", 40: "cfo",
            43: "revolver_flow", 33: "balance_check"}
    for r, key in keys.items():
        CELLS[f"stmt_{key}_row"] = f"Statements!C{r}:N{r}"
    ws.freeze_panes = "B5"
    widths(ws, {"A": 34, **{get_column_letter(c): 12.5 for c in range(2, 16)}})


def reforecast(wb) -> None:
    ws = wb.create_sheet("Reforecast")
    title(ws, "Reforecast 3+9: Q1 actuals, then nine months re-driven",
          "Apr-Dec take Q1's volume run-rate by channel, unit cost by category and opex by department; prices stay at budget.")
    b = f"Build!${{c}}${B0}:${{c}}${B1}"
    q = f"Data_Q1!${{c}}${Q0}:${{c}}${Q1R}"
    qo = f"Data_Q1!${{c}}${QO0}:${{c}}${QO1}"
    header(ws, 5, ["Channel", "Q1 budget cases", "Q1 actual cases", "Volume ratio"])
    for k, ch in enumerate(A["channels"]):
        r = 6 + k
        ws.cell(row=r, column=1, value=ch)
        put(ws, f"B{r}", f'=SUMIFS({b.format(c="F")},{b.format(c="D")},A{r},{b.format(c="A")},"<=3")', NUM1)
        put(ws, f"C{r}", f'=SUMIFS({q.format(c="D")},{q.format(c="C")},A{r})', NUM1)
        put(ws, f"D{r}", f"=C{r}/B{r}", "0.0000", f"vol_ratio_{k}")
    header(ws, 11, ["Category", "Q1 budget cost / case", "Q1 actual cost / case", "Cost ratio"])
    for k, cat in enumerate(A["categories"]):
        r = 12 + k
        ws.cell(row=r, column=1, value=cat)
        put(ws, f"B{r}", f'=SUMIFS({b.format(c="M")},{b.format(c="C")},A{r},{b.format(c="A")},"<=3")'
                         f'/SUMIFS({b.format(c="F")},{b.format(c="C")},A{r},{b.format(c="A")},"<=3")', MONEY2)
        put(ws, f"C{r}", f'=SUMIFS({q.format(c="F")},{q.format(c="B")},A{r})/SUMIFS({q.format(c="D")},{q.format(c="B")},A{r})',
            MONEY2)
        put(ws, f"D{r}", f"=C{r}/B{r}", "0.0000", f"cost_ratio_{k}")
    header(ws, 19, ["Department", "Q1 budget opex", "Q1 actual opex", "Opex ratio", "Apr-Dec budget", "Apr-Dec reforecast"])
    for k, d in enumerate(A["departments"]):
        r = 20 + k
        ws.cell(row=r, column=1, value=d)
        put(ws, f"B{r}", f"=SUMPRODUCT((Opex!$A$13:$A$16=A{r})*(Opex!$C$5:$N$5<=3)*Opex!$C$13:$N$16)", MONEY)
        put(ws, f"C{r}", f'=SUMIFS({qo.format(c="C")},{qo.format(c="B")},A{r})', MONEY)
        put(ws, f"D{r}", f"=C{r}/B{r}", "0.0000", f"opex_ratio_{k}")
        put(ws, f"E{r}", f"=SUMPRODUCT((Opex!$A$13:$A$16=A{r})*(Opex!$C$5:$N$5>3)*Opex!$C$13:$N$16)", MONEY)
        put(ws, f"F{r}", f"=E{r}*D{r}", MONEY)

    header(ws, 26, ["FY2027", "Budget (V2)", "Reforecast 3+9", "Variance"])
    ws["A27"], ws["A28"], ws["A29"], ws["A30"] = "Net revenue", "Gross margin", "Opex", "EBITDA"
    put(ws, "B27", "=Consolidation!E6", MONEY)
    put(ws, "C27", f"=SUM({b.format(c='Q')})", MONEY, "rf_revenue")
    put(ws, "B28", "=Consolidation!E9", MONEY)
    put(ws, "C28", f"=SUM({b.format(c='Q')})-SUM({b.format(c='R')})-SUM({b.format(c='S')})", MONEY, "rf_gross_margin")
    put(ws, "B29", "=Consolidation!E10", MONEY)
    put(ws, "C29", f"=SUM({qo.format(c='C')})+SUM(F20:F23)", MONEY, "rf_opex")
    put(ws, "B30", "=B28-B29", MONEY, bold=True)
    put(ws, "C30", "=C28-C29", MONEY, "rf_ebitda", bold=True)
    for r in range(27, 31):
        put(ws, f"D{r}", f"=C{r}-B{r}", MONEY)

    ws["A32"] = "EBITDA bridge: budget to reforecast"
    ws["A32"].font = SECTION_FONT
    ws["A33"] = "Q1: actual against budget"
    put(ws, "B33",
        f'=(SUMIFS({b.format(c="Q")},{b.format(c="A")},"<=3")-SUMIFS({b.format(c="R")},{b.format(c="A")},"<=3")'
        f'-SUMIFS({b.format(c="S")},{b.format(c="A")},"<=3")-SUM(C20:C23))'
        f'-(SUMIFS({b.format(c="K")},{b.format(c="A")},"<=3")-SUMIFS({b.format(c="M")},{b.format(c="A")},"<=3")'
        f'-SUMIFS({b.format(c="O")},{b.format(c="A")},"<=3")-SUM(B20:B23))', MONEY, "bridge_rf_q1")
    ws["A34"] = "Apr-Dec: volume, at budget margin per case"
    put(ws, "B34",
        f"=SUMPRODUCT(({b.format(c='A')}>3)*({b.format(c='P')}-{b.format(c='F')})"
        f"*({b.format(c='G')}*(1+{b.format(c='H')})*(1+{b.format(c='I')})-{b.format(c='L')}-{b.format(c='N')}))",
        MONEY, "bridge_rf_volume")
    ws["A35"] = "Apr-Dec: unit cost"
    put(ws, "B35", f"=-SUMPRODUCT(({b.format(c='A')}>3)*({b.format(c='R')}-{b.format(c='P')}*{b.format(c='L')}))",
        MONEY, "bridge_rf_cost")
    ws["A36"] = "Apr-Dec: opex run-rate"
    put(ws, "B36", "=-(SUM(F20:F23)-SUM(E20:E23))", MONEY, "bridge_rf_opex")
    ws["A37"] = "Budget to reforecast, total"
    put(ws, "B37", "=SUM(B33:B36)", MONEY, "bridge_rf_total", bold=True)
    widths(ws, {"A": 42, "B": 18, "C": 18, "D": 14, "E": 16, "F": 18})


def checks(wb) -> None:
    ws = wb.create_sheet("Checks")
    title(ws, "Checks: the workbook proves itself", "Each check states its tolerance. Status is a formula.")
    header(ws, 4, ["Check", "Difference", "Tolerance", "Status"])
    b = f"Build!${{c}}${B0}:${{c}}${B1}"
    rows = [
        ("Seasonality: each category's twelve shares sum to one",
         "=" + "+".join(f"ABS(SUM(Data_Base!B{r}:M{r})-1)" for r in range(32, 38)), 1e-9),
        ("Build: cases = base cases x (1 + channel growth), in total",
         f"=SUM({b.format(c='F')})-SUMPRODUCT(SUMIFS(Data_Base!$C$5:$C$28,Data_Base!$B$5:$B$28,Assumptions!$A$19:$A$22),"
         "1+Assumptions!$B$19:$B$22)", 1e-6),
        ("Build: V1 revenue before the price change = base revenue share",
         f'=SUMIFS({b.format(c="J")},{b.format(c="A")},"<"&Assumptions!C37)-SUMPRODUCT(({b.format(c="A")}<Assumptions!C37)'
         f'*{b.format(c="F")}*{b.format(c="G")})', 0.01),
        ("Statements: revenue = consolidated V2 revenue", "=Statements!O6-Consolidation!E6", 0.01),
        ("Statements: opex = consolidated V2 opex", "=Statements!O10-Consolidation!E10", 0.01),
        ("Statements: EBITDA = consolidated V2 EBITDA", "=Statements!O11-Consolidation!E11", 0.01),
        ("Balance sheet balances, every month", "=SUMPRODUCT(ABS(Statements!B33:N33))", 0.01),
        ("Cash flow ties to the balance sheet, every month", "=SUMPRODUCT(ABS(Statements!C45:N45))", 0.01),
        ("Revolver never negative", "=-MIN(0,MIN(Statements!C30:N30))", 0.01),
        ("Cash never below the minimum", "=MAX(0,Assumptions!C42-MIN(Statements!C23:N23))", 0.01),
        ("Equity moves only by net income", "=Statements!O31-Statements!B31-Statements!O17", 0.01),
        ("Challenge bridge closes: V1 + levers = V2 EBITDA", "=Consolidation!D11+Consolidation!B24-Consolidation!E11", 0.01),
        ("Reforecast Q1 = actuals (revenue)", f'=SUMIFS({b.format(c="Q")},{b.format(c="A")},"<=3")-SUM(Data_Q1!E{Q0}:E{Q1R})', 0.01),
        ("Reforecast Q1 = actuals (opex)", "=SUM(Reforecast!C20:C23)-SUM(Data_Q1!C80:C91)", 0.01),
        ("Reforecast bridge closes", "=Reforecast!B37-Reforecast!D30", 0.01),
    ]
    for k, (label, formula, tol) in enumerate(rows):
        r = 5 + k
        ws.cell(row=r, column=1, value=label)
        put(ws, f"B{r}", formula, "0.000000")
        ws.cell(row=r, column=3, value=tol).number_format = "0.0E+00"
        put(ws, f"D{r}", f'=IF(ABS(B{r})<=C{r},"PASS","FAIL")')
    last = 5 + len(rows) - 1
    ws[f"A{last + 2}"] = "Checks passing"
    put(ws, f"B{last + 2}", f'=COUNTIF(D5:D{last},"PASS")', "0", "checks_passing", bold=True)
    ws[f"C{last + 2}"] = f"of {len(rows)}"
    CELLS["checks_total"] = str(len(rows))
    CELLS["checks_range"] = f"Checks!D5:D{last}"
    widths(ws, {"A": 60, "B": 16, "C": 12, "D": 10})


def summary(wb) -> None:
    ws = wb.active
    ws.title = "Summary"
    title(ws, "Kestrel Bay Provisions: FY2027 budget cycle", "Synthetic BC food distributor. Every figure is a formula.")
    ws["A4"] = "The cycle"
    ws["A4"].font = SECTION_FONT
    steps = [
        ("V0 targets", "Leadership sets revenue growth and an EBITDA margin on the last twelve months."),
        ("V1 submissions", "Sales submit channel volumes and category prices; procurement, cost inflation; "
                           "logistics, freight; department heads, opex and hires. Phased on 2025's seasonality."),
        ("Consolidation", "Bottom-up against top-down: the gap, in revenue and EBITDA."),
        ("V2 challenge", "Foodservice price on beef and seafood, one hire deferred, warehouse non-pay cut, "
                         "technology growth capped. V2 is the approved budget."),
        ("Three statements", "Monthly P&L, balance sheet and cash flow for the approved budget, with a revolver."),
        ("3+9 reforecast", "After Q1: actuals, then nine months re-driven by Q1's run-rates, and the bridge."),
    ]
    for k, (a, t) in enumerate(steps):
        ws.cell(row=5 + k, column=1, value=a).font = TOTAL_FONT
        ws.cell(row=5 + k, column=2, value=t)
    header(ws, 12, ["FY2027", "Net revenue", "EBITDA", "EBITDA margin"])
    lines = [("Last twelve months (base)", "B"), ("V0 target", "C"), ("V1 bottom-up", "D"), ("V2 approved", "E")]
    for k, (label, col) in enumerate(lines):
        r = 13 + k
        ws.cell(row=r, column=1, value=label)
        put(ws, f"B{r}", f"=Consolidation!{col}6", MONEY)
        put(ws, f"C{r}", f"=Consolidation!{col}11", MONEY)
        put(ws, f"D{r}", f"=Consolidation!{col}12", PCT2)
    ws["A17"] = "Reforecast 3+9"
    put(ws, "B17", "=Reforecast!C27", MONEY)
    put(ws, "C17", "=Reforecast!C30", MONEY)
    put(ws, "D17", "=C17/B17", PCT2)
    ws["A19"] = "Approved EBITDA still short of target by"
    put(ws, "B19", "=-Consolidation!E16", MONEY, bold=True)
    ws["A20"] = "Further price needed from the effective month to close it"
    put(ws, "B20", "=Consolidation!B27", PCT2, bold=True)
    ws["A21"] = "Net income, approved budget"
    put(ws, "B21", "=Statements!O17", MONEY)
    ws["A22"] = "Closing revolver / cash"
    put(ws, "B22", "=Statements!O30", MONEY)
    put(ws, "C22", "=Statements!O23", MONEY)
    ws["A23"] = "Checks passing"
    put(ws, "B23", "=Checks!B21", "0")
    ws["C23"] = "of 15"
    widths(ws, {"A": 52, "B": 96, "C": 16, "D": 14})
    for r in range(12, 24):
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="right")


def recalculate(path: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "out"
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                       check=True, capture_output=True, timeout=300)
        shutil.copyfile(out / path.name, path)


def main() -> int:
    wb = Workbook()
    summary(wb)
    assumptions(wb)
    consolidation(wb)
    statements(wb)
    reforecast(wb)
    build_sheet(wb)
    opex_sheet(wb)
    data_sheets(wb)
    checks(wb)
    for name, ref in [("v2_ebitda", "Consolidation!$E$11"), ("v2_revenue", "Consolidation!$E$6"),
                      ("rf_ebitda", "Reforecast!$C$30"), ("checks_passing", "Checks!$B$21")]:
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    wb.save(OUT)
    recalculate(OUT)
    (HERE / "cell_map.json").write_text(json.dumps(CELLS, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {OUT.name}: {len(CELLS)} published cells mapped")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
